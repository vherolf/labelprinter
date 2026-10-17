"""Read a CSV or Excel file into rows for batch printing. The first row holds the column names."""

import csv
import datetime
import io

from openpyxl import load_workbook


class TableError(Exception):
    pass


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d" if value.time() == datetime.time() else "%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value).replace("\r", "").strip()


def _read_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    # Excel exports use ";" in many locales: pick the separator the header line uses most
    header = text.split("\n", 1)[0]
    delimiter = max(",;\t", key=header.count)
    return [[_cell(v) for v in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises a zoo of exception types on bad files
        raise TableError(f"Can't read Excel file: {e}")
    try:
        return [[_cell(v) for v in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def parse(filename: str, data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        raw = _read_xlsx(data)
    elif name.endswith((".csv", ".tsv", ".txt")):
        raw = _read_csv(data)
    else:
        raise TableError("Unsupported file type - use .csv or .xlsx")

    raw = [r for r in raw if any(r)]
    if not raw:
        raise TableError("The file is empty")
    header = raw[0]
    while header and not header[-1]:
        header.pop()
    columns = [h or f"Column {i + 1}" for i, h in enumerate(header)]
    if len(set(columns)) != len(columns):
        raise TableError("Column names in the first row must be unique")
    rows = [dict(zip(columns, r + [""] * (len(columns) - len(r)))) for r in raw[1:]]
    if not rows:
        raise TableError("No data rows below the header row")
    return columns, rows
